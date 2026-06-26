import os
import sys
import json
import time
import logging
from pathlib import Path
from typing import List, Dict, Any

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

def parse_pdf_with_docling(file_path: str) -> Dict[str, Any]:
    try:
        from docling.backend.docling_parse_v2_backend import DoclingParseV2DocumentBackend
        from docling.document_converter import DocumentConverter, FormatOption
        from docling.datamodel.pipeline_options import PdfPipelineOptions, TableFormerMode, EasyOcrOptions
        from docling.datamodel.base_models import InputFormat
        from docling.pipeline.standard_pdf_pipeline import StandardPdfPipeline
        from docling.datamodel.base_models import ConversionStatus
        
        pipeline_options = PdfPipelineOptions()
        pipeline_options.do_ocr = True
        ocr_options = EasyOcrOptions(lang=['zh', 'en'], force_full_page_ocr=False)
        pipeline_options.ocr_options = ocr_options
        pipeline_options.do_table_structure = True
        pipeline_options.table_structure_options.do_cell_matching = True
        pipeline_options.table_structure_options.mode = TableFormerMode.ACCURATE
        
        format_options = {
            InputFormat.PDF: FormatOption(
                pipeline_cls=StandardPdfPipeline,
                pipeline_options=pipeline_options,
                backend=DoclingParseV2DocumentBackend
            )
        }
        
        doc_converter = DocumentConverter(format_options=format_options)
        conv_results = doc_converter.convert_all(source=[Path(file_path)])
        
        for conv_res in conv_results:
            if conv_res.status == ConversionStatus.SUCCESS:
                data = conv_res.document.export_to_dict()
                return assemble_report(data)
            else:
                logger.error(f"Failed to convert PDF: {conv_res.input.file}")
                return {"error": f"Failed to convert PDF: {conv_res.input.file}"}
    except ImportError as e:
        logger.error(f"Docling not installed: {e}")
        return {"error": f"Docling not installed: {e}"}
    except Exception as e:
        logger.error(f"Error parsing PDF: {e}", exc_info=True)
        return {"error": f"Error parsing PDF: {str(e)}"}

def assemble_report(data: Dict[str, Any]) -> Dict[str, Any]:
    assembled = {}
    assembled['metainfo'] = assemble_metainfo(data)
    assembled['content'] = assemble_content(data)
    assembled['tables'] = assemble_tables(data)
    assembled['text'] = extract_full_text(data)
    return assembled

def assemble_metainfo(data: Dict[str, Any]) -> Dict[str, Any]:
    metainfo = {}
    origin = data.get('origin', {})
    filename = origin.get('filename', '')
    metainfo['filename'] = filename
    metainfo['sha1_name'] = filename.rsplit('.', 1)[0] if '.' in filename else filename
    metainfo['pages_amount'] = len(data.get('pages', []))
    metainfo['text_blocks_amount'] = len(data.get('texts', []))
    metainfo['tables_amount'] = len(data.get('tables', []))
    metainfo['pictures_amount'] = len(data.get('pictures', []))
    return metainfo

def extract_full_text(data: Dict[str, Any]) -> str:
    texts = data.get('texts', [])
    text_content = []
    for text_item in texts:
        text = text_item.get('text', '')
        if text:
            text_content.append(text)
    return '\n'.join(text_content)

def assemble_content(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    pages = {}
    body_children = data.get('body', {}).get('children', [])
    groups = data.get('groups', [])
    
    expanded_children = expand_groups(body_children, groups)
    
    for item in expanded_children:
        if isinstance(item, dict) and '$ref' in item:
            ref = item['$ref']
            parts = ref.split('/')[-2:]
            if len(parts) == 2:
                ref_type, ref_num = parts
                ref_num = int(ref_num)
                
                if ref_type == 'texts':
                    text_item = data.get('texts', [])[ref_num]
                    content_item = {
                        'text': text_item.get('text', ''),
                        'type': text_item.get('label', 'text'),
                        'text_id': ref_num
                    }
                    
                    if 'group_id' in item:
                        content_item['group_id'] = item['group_id']
                        content_item['group_name'] = item.get('group_name', '')
                        content_item['group_label'] = item.get('group_label', '')
                    
                    if 'prov' in text_item and text_item['prov']:
                        page_num = text_item['prov'][0]['page_no']
                        if page_num not in pages:
                            pages[page_num] = {'page': page_num, 'content': []}
                        pages[page_num]['content'].append(content_item)
    
    return [pages[page_num] for page_num in sorted(pages.keys())]

def expand_groups(body_children: List[Any], groups: List[Dict[str, Any]]) -> List[Any]:
    expanded = []
    for item in body_children:
        if isinstance(item, dict) and '$ref' in item:
            ref = item['$ref']
            parts = ref.split('/')[-2:]
            if len(parts) == 2:
                ref_type, ref_num = parts
                ref_num = int(ref_num)
                
                if ref_type == 'groups':
                    group = groups[ref_num]
                    group_id = ref_num
                    group_name = group.get('name', '')
                    group_label = group.get('label', '')
                    
                    for child in group.get('children', []):
                        child_copy = child.copy()
                        child_copy['group_id'] = group_id
                        child_copy['group_name'] = group_name
                        child_copy['group_label'] = group_label
                        expanded.append(child_copy)
                else:
                    expanded.append(item)
        else:
            expanded.append(item)
    return expanded

def assemble_tables(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    assembled = []
    tables = data.get('tables', [])
    
    for i, table in enumerate(tables):
        table_json_obj = table
        table_md = table_to_md(table_json_obj)
        
        page_num = 1
        bbox = []
        if 'prov' in table and table['prov']:
            page_num = table['prov'][0]['page_no']
            bbox_data = table['prov'][0].get('bbox', {})
            bbox = [bbox_data.get('l', 0), bbox_data.get('t', 0), 
                    bbox_data.get('r', 0), bbox_data.get('b', 0)]
        
        nrows = table.get('data', {}).get('num_rows', 0)
        ncols = table.get('data', {}).get('num_cols', 0)
        
        table_obj = {
            'table_id': i,
            'page': page_num,
            'bbox': bbox,
            '#-rows': nrows,
            '#-cols': ncols,
            'markdown': table_md,
            'json': table_json_obj
        }
        assembled.append(table_obj)
    
    return assembled

def table_to_md(table: Dict[str, Any]) -> str:
    from tabulate import tabulate
    
    table_data = []
    grid = table.get('data', {}).get('grid', [])
    for row in grid:
        table_row = [cell.get('text', '') for cell in row]
        table_data.append(table_row)
    
    if len(table_data) > 1 and len(table_data[0]) > 0:
        try:
            md_table = tabulate(table_data[1:], headers=table_data[0], tablefmt="github")
        except ValueError:
            md_table = tabulate(table_data[1:], headers=table_data[0], tablefmt="github", disable_numparse=True)
    else:
        md_table = tabulate(table_data, tablefmt="github")
    
    return md_table

def parse_word(file_path: str) -> Dict[str, Any]:
    try:
        from docx import Document
        
        doc = Document(file_path)
        text_content = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)
        
        tables_md = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                from tabulate import tabulate
                md_table = tabulate(table_data, tablefmt="github")
                tables_md.append(md_table)
        
        return {
            'metainfo': {
                'filename': os.path.basename(file_path),
                'sha1_name': os.path.basename(file_path).rsplit('.', 1)[0],
                'pages_amount': 1,
                'text_blocks_amount': len(text_content),
                'tables_amount': len(tables_md),
                'pictures_amount': 0
            },
            'content': [{
                'page': 1,
                'content': [{'text': text, 'type': 'paragraph'} for text in text_content]
            }],
            'tables': tables_md,
            'text': '\n'.join(text_content)
        }
    except Exception as e:
        logger.error(f"Error parsing Word: {e}", exc_info=True)
        return {"error": f"Error parsing Word: {str(e)}"}

def parse_excel(file_path: str) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook
        
        wb = load_workbook(file_path)
        all_tables_md = []
        all_text = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            table_data = []
            for row in sheet.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else '' for cell in row]
                if any(row_data):
                    table_data.append(row_data)
            
            if table_data:
                from tabulate import tabulate
                md_table = f"## {sheet_name}\n\n" + tabulate(table_data, tablefmt="github")
                all_tables_md.append(md_table)
                all_text.append(md_table)
        
        return {
            'metainfo': {
                'filename': os.path.basename(file_path),
                'sha1_name': os.path.basename(file_path).rsplit('.', 1)[0],
                'pages_amount': len(wb.sheetnames),
                'text_blocks_amount': len(all_tables_md),
                'tables_amount': len(all_tables_md),
                'pictures_amount': 0
            },
            'content': [{
                'page': i + 1,
                'content': [{'text': table, 'type': 'table'}]
            } for i, table in enumerate(all_tables_md)],
            'tables': all_tables_md,
            'text': '\n\n'.join(all_text)
        }
    except Exception as e:
        logger.error(f"Error parsing Excel: {e}", exc_info=True)
        return {"error": f"Error parsing Excel: {str(e)}"}

def parse_text(file_path: str) -> Dict[str, Any]:
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        lines = content.split('\n')
        text_blocks = [line for line in lines if line.strip()]
        
        return {
            'metainfo': {
                'filename': os.path.basename(file_path),
                'sha1_name': os.path.basename(file_path).rsplit('.', 1)[0],
                'pages_amount': 1,
                'text_blocks_amount': len(text_blocks),
                'tables_amount': 0,
                'pictures_amount': 0
            },
            'content': [{
                'page': 1,
                'content': [{'text': line, 'type': 'paragraph'} for line in text_blocks]
            }],
            'tables': [],
            'text': content
        }
    except Exception as e:
        logger.error(f"Error parsing text: {e}", exc_info=True)
        return {"error": f"Error parsing text: {str(e)}"}

def parse_markdown(file_path: str) -> Dict[str, Any]:
    return parse_text(file_path)

def parse_document(file_path: str) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return parse_pdf_with_docling(file_path)
    elif ext in ['.docx', '.doc']:
        return parse_word(file_path)
    elif ext in ['.xlsx', '.xls']:
        return parse_excel(file_path)
    elif ext == '.txt':
        return parse_text(file_path)
    elif ext == '.md':
        return parse_markdown(file_path)
    else:
        logger.error(f"Unsupported file type: {ext}")
        return {"error": f"Unsupported file type: {ext}"}

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Please provide file path as argument"}, ensure_ascii=False))
        sys.exit(1)
    
    file_path = sys.argv[1]
    result = parse_document(file_path)
    print(json.dumps(result, ensure_ascii=False))
